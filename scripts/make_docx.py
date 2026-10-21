import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(ROOT, 'docs', 'eqs')
os.makedirs(OUTDIR, exist_ok=True)

plt.rcParams['mathtext.fontset'] = 'cm'

def render_eq(latex, filename, fontsize=22):
    fig = plt.figure(figsize=(7, 1.1))
    fig.text(0.5, 0.5, latex, ha='center', va='center', fontsize=fontsize)
    path = os.path.join(OUTDIR, filename)
    fig.savefig(path, dpi=200, bbox_inches='tight', transparent=False, pad_inches=0.1,
                facecolor='white')
    plt.close(fig)
    return path

eq_balance = render_eq(
    r'$\rho\, c_p\, \dfrac{\partial T}{\partial t} = k\, \dfrac{\partial^2 T}{\partial x^2} + \dot{q}$',
    'eq_balance.png')
eq_alpha = render_eq(
    r'$\dfrac{\partial T}{\partial t} = \alpha\, \dfrac{\partial^2 T}{\partial x^2} + \dfrac{\dot{q}}{\rho\, c_p}, \qquad \alpha = \dfrac{k}{\rho\, c_p}$',
    'eq_alpha.png')
eq_ftcs = render_eq(
    r'$T_i^{\,n+1} = T_i^{\,n} + Fo\left(T_{i+1}^{\,n} - 2\,T_i^{\,n} + T_{i-1}^{\,n}\right)$',
    'eq_ftcs.png')
eq_fourier = render_eq(
    r'$Fo = \dfrac{\alpha\, \Delta t}{(\Delta x)^2}$',
    'eq_fourier.png')
eq_stab = render_eq(
    r'$Fo \leq \dfrac{1}{2}$',
    'eq_stab.png')
eq_ci = render_eq(r'$T(x,0) = T_0(x)$', 'eq_ci.png', fontsize=20)
eq_dir = render_eq(r'$T(x_b,\, t) = T_{\mathrm{fija}}$', 'eq_dir.png', fontsize=20)
eq_neu = render_eq(r'$-k\, \left.\dfrac{\partial T}{\partial x}\right|_{x_b} = q\'\'(t)$', 'eq_neu.png', fontsize=20)
eq_rob = render_eq(r'$-k\, \left.\dfrac{\partial T}{\partial x}\right|_{x_b} = h\left(T_s(t) - T_\infty(t)\right)$', 'eq_rob.png', fontsize=20)

doc = Document()

doc.add_heading('Respuestas E01 - BIM, Difusion Termica y Solver 1D', 0)
doc.add_paragraph('Preparacion revision oral - Semana 1 / miercoles')
doc.add_paragraph('')

doc.add_heading('1) BIM / IFC', level=1)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Pregunta'
hdr[1].text = 'Respuesta'
bim_data = [
    ('Que es BIM?',
     'BIM (Building Information Modeling) es un proceso de modelado digital que integra '
     'geometria e informacion del proyecto (materiales, propiedades, elementos, secuencia '
     'constructiva). No es solo un dibujo 3D: es una base de datos del elemento constructivo '
     'que conecta geometria con datos tecnicos.'),
    ('Que es IFC?',
     'IFC (Industry Foundation Classes) es un formato abierto (.ifc) estandarizado por '
     'buildingSMART (ISO 16739). Sirve para intercambiar informacion BIM entre distintos '
     'programas sin depender de un software propietario.'),
    ('Que rol cumple Blender?',
     'Blender es un modelador 3D de codigo abierto. Es la interfaz grafica para crear, editar '
     'y visualizar la geometria 3D de la fundacion. Por si solo no maneja la semantica BIM/IFC.'),
    ('Que agrega Bonsai?',
     'Bonsai (ex BlenderBIM) es un add-on de Blender que agrega la capa BIM/IFC. Permite '
     'clasificar elementos (IfcFooting), crear la jerarquia espacial (IfcProject, IfcSite, '
     'IfcBuilding, IfcBuildingStorey), asignar Property Sets y exportar un .ifc con semantica.'),
    ('Por que el solver no depende de Blender?',
     'Porque la fuente de verdad es el archivo .ifc, no el .blend. El flujo es: '
     'Blender + Bonsai -> IFC -> IfcOpenShell -> geometria + propiedades -> voxelizacion -> solver. '
     'Al leer el IFC con IfcOpenShell fuera de Blender, el solver corre por linea de comandos, '
     'en tests, en HPC y es reproducible. Si importaramos bpy, solo funcionaria dentro de Blender.'),
]
for q, a in bim_data:
    row = table.add_row().cells
    row[0].text = q
    row[1].text = a

doc.add_heading('2) Fisica de difusion termica', level=1)
doc.add_paragraph(
    'La difusion de calor describe como se propaga la temperatura en un material por conduccion, '
    'y como el calor generado internamente (hidratacion del cemento) eleva la temperatura. '
    'Las propiedades fisicas clave son:'
)
props = doc.add_table(rows=1, cols=3)
props.style = 'Light Grid Accent 1'
ph = props.rows[0].cells
ph[0].text = 'Variable'
ph[1].text = 'Significado fisico'
ph[2].text = 'Unidades (SI)'
props_data = [
    ('k', 'Conductividad termica. Que tan facil se transfiere calor por conduccion.', 'W/(m K)'),
    ('rho', 'Densidad. Masa por unidad de volumen.', 'kg/m3'),
    ('c_p', 'Calor especifico a presion constante. Energia para subir 1 kg en 1 K.', 'J/(kg K)'),
    ('alpha', 'Difusividad termica. alpha = k/(rho c_p). Rapidez de propagacion del cambio de temperatura.', 'm2/s'),
]
for v, s, u in props_data:
    row = props.add_row().cells
    row[0].text = v
    row[1].text = s
    row[2].text = u

doc.add_heading('Balance de energia', level=2)
doc.add_paragraph(
    'La variacion de energia interna dentro de un volumen de control es igual al calor neto '
    'que entra y sale por conduccion, mas el calor generado internamente. Esto da la ecuacion '
    'de difusion con termino fuente (generacion volumetrica q_dot):'
)
doc.add_picture(eq_balance, width=Inches(3.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Dividiendo por rho*c_p se obtiene la forma en funcion de la difusividad:')
doc.add_picture(eq_alpha, width=Inches(4.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('3) Condicion inicial y condiciones de borde', level=1)

doc.add_paragraph('Condicion inicial (CI). Define la temperatura de todo el dominio al inicio '
                  'de la simulacion. En el proyecto suele ser la temperatura de colocacion:')
doc.add_picture(eq_ci, width=Inches(1.8))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Dirichlet. Se impone una temperatura fija directamente en el borde:')
doc.add_picture(eq_dir, width=Inches(2.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Neumann. Se impone un flujo de calor. Si q\'\' = 0 el borde es adiabatico '
                  '(perfectamente aislado, sin intercambio de calor):')
doc.add_picture(eq_neu, width=Inches(2.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Robin. Representa conveccion entre la superficie (T_s) y el ambiente (T_inf). '
                  'Es la condicion mas usada para el intercambio con el aire:')
doc.add_picture(eq_rob, width=Inches(3.6))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('4) Discretizacion 1D y estabilidad', level=1)
doc.add_paragraph(
    'Se usa el esquema explicito FTCS (Forward Time, Centered Space). La derivada temporal se '
    'aproxima hacia adelante y la espacial de forma centrada. La temperatura en cada nodo i se '
    'actualiza con:'
)
doc.add_picture(eq_ftcs, width=Inches(4.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('donde el numero de Fourier es:')
doc.add_picture(eq_fourier, width=Inches(2.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('Estabilidad del esquema explicito', level=2)
doc.add_paragraph(
    'El esquema explicito es condicionalmente estable. Para conduccion pura 1D debe cumplirse '
    'la condicion de Von Neumann:'
)
doc.add_picture(eq_stab, width=Inches(1.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(
    'Si Fo supera 0.5 la solucion se vuelve inestable y aparecen oscilaciones numericas que '
    'crecen exponencialmente. Por eso dt se elige lo suficientemente pequeno para que Fo <= 0.5, '
    'o bien se usa un esquema implicito.'
)

doc.add_heading('Checklist para la revision', level=1)
for item in [
    'URL del repositorio y commit (SHA) evaluado.',
    'modelo/foundation.ifc con las 5 vertidas (IfcFooting) y su Property Set.',
    'Saber explicar el solver 1D: loop temporal, calculo de T_i^{n+1}, condiciones de borde y Fo.',
    'Repositorio reproducible: README minimo, requirements.txt y primeros tests.',
    'Capacidad de navegar por el codigo real durante la revision.',
]:
    doc.add_paragraph(item, style='List Bullet')

out = os.path.join(ROOT, 'docs', 'E01_respuestas_bim_difusion_solver1d.docx')
doc.save(out)
print('DONE', out)
